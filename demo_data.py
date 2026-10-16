"""Create a repeatable, entirely fictional Excel input for the portfolio demo."""

from datetime import date
from pathlib import Path
import random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


COLUMNS = [
    "batch_id", "date", "client", "base_kg", "pieces", "material", "category",
    "unit", "quantity", "unit_price", "line_cost",
]
MATERIALS = [
    ("Material A", "Materials", "kg", 0.055, 28.0),
    ("Material B", "Materials", "kg", 0.035, 20.0),
    ("Material C", "Materials", "kg", 0.020, 18.0),
    ("Water", "Utilities", "L", 5.0, 0.025),
    ("Electricity", "Utilities", "kWh", 0.14, 2.5),
    ("Packaging", "Other", "units", 0.02, 5.0),
]


def generate_workbook(path: Path, seed: int = 42) -> Path:
    """Write invented records without reading any source company files."""
    rng = random.Random(seed)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(COLUMNS)
    batch_number = 0
    for month in range(5, 9):
        for client_index, client in enumerate(["Client A", "Client B", "Client C"]):
            for day in [3, 7, 12, 17, 22, 27]:
                batch_number += 1
                base_kg = round(rng.uniform(2400, 4100), 1)
                pieces = round(base_kg / rng.uniform(23, 29))
                for material, category, unit, rate, price in MATERIALS:
                    usage_factor = 1 + client_index * 0.04 + rng.uniform(-0.09, 0.09)
                    quantity = round(base_kg * rate * usage_factor, 2)
                    unit_price = round(price * (1 + (month - 5) * 0.015), 3)
                    sheet.append([
                        f"DEMO-{batch_number:03d}", date(2026, month, day), client,
                        base_kg, pieces, material, category, unit, quantity,
                        unit_price, round(quantity * unit_price, 2),
                    ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 18
    for row in sheet.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
    notes = workbook.create_sheet("About")
    notes.append(["Portfolio demo — synthetic data"])
    notes.append(["Every client, batch, price, quantity and pattern is fictional."])
    notes.append(["Generated from a fixed random seed; no private records are read."])
    notes.append(["Currency is illustrative USD and does not represent business results."])
    notes.column_dimensions["A"].width = 90
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


if __name__ == "__main__":
    output = generate_workbook(Path(__file__).resolve().parent / "demo" / "synthetic_records.xlsx")
    print(f"Created synthetic input: {output}")
