"""Excel parsing, aggregation, Pareto analysis and export for synthetic records."""

from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from demo_data import COLUMNS


def read_records(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook["Records"].iter_rows(values_only=True)
        headers = list(next(values))
        if headers != COLUMNS:
            raise ValueError("Unexpected demo workbook columns.")
        frame = pd.DataFrame(values, columns=headers).dropna(how="all")
    finally:
        workbook.close()
    for column in ["batch_id", "client", "material", "category", "unit"]:
        frame[column] = frame[column].astype(str).str.strip()
    frame["date"] = pd.to_datetime(frame["date"], errors="raise")
    for column in ["base_kg", "pieces", "quantity", "unit_price", "line_cost"]:
        frame[column] = pd.to_numeric(frame[column], errors="raise")
    if frame.isna().any().any() or (frame[["base_kg", "pieces"]] <= 0).any().any():
        raise ValueError("Demo records must be complete, with positive batch sizes.")
    if (frame[["quantity", "unit_price", "line_cost"]] < 0).any().any():
        raise ValueError("Demo records cannot have negative costs or quantities.")
    calculated = (frame["quantity"] * frame["unit_price"]).round(2)
    if ((calculated - frame["line_cost"]).abs() > 0.011).any():
        raise ValueError("Line costs do not match quantity times unit price.")
    batch_fields = ["date", "client", "base_kg", "pieces"]
    if (frame.groupby("batch_id")[batch_fields].nunique() > 1).any().any():
        raise ValueError("A batch has conflicting identifying fields or sizes.")
    frame["month"] = frame["date"].dt.strftime("%Y-%m")
    return frame


def filter_records(frame: pd.DataFrame, clients=None, months=None) -> pd.DataFrame:
    selected = frame
    if clients:
        selected = selected[selected["client"].isin(clients)]
    if months:
        selected = selected[selected["month"].isin(months)]
    return selected.copy()


def batch_summary(frame: pd.DataFrame) -> pd.DataFrame:
    batches = frame.groupby("batch_id", as_index=False, sort=True).agg(
        date=("date", "first"), client=("client", "first"),
        month=("month", "first"), base_kg=("base_kg", "first"),
        pieces=("pieces", "first"), cost=("line_cost", "sum"),
    )
    batches["cost_per_kg"] = batches["cost"] / batches["base_kg"]
    return batches


def metrics(frame: pd.DataFrame) -> dict:
    batches = batch_summary(frame)
    cost = float(batches["cost"].sum())
    kilograms = float(batches["base_kg"].sum())
    return {
        "total_cost": cost, "base_kg": kilograms,
        "cost_per_kg": cost / kilograms if kilograms else 0.0,
        "batches": len(batches), "clients": int(batches["client"].nunique()),
    }


def monthly_summary(frame: pd.DataFrame) -> pd.DataFrame:
    summary = batch_summary(frame).groupby(["month", "client"], as_index=False).agg(
        batches=("batch_id", "count"), base_kg=("base_kg", "sum"), cost=("cost", "sum"),
    )
    summary["cost_per_kg"] = summary["cost"] / summary["base_kg"]
    return summary


def material_pareto(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.groupby(["material", "category", "unit"], as_index=False).agg(
        quantity=("quantity", "sum"), cost=("line_cost", "sum"),
    ).sort_values(["cost", "material"], ascending=[False, True]).reset_index(drop=True)
    total = result["cost"].sum()
    result["share"] = result["cost"] / total if total else 0.0
    result["cumulative_share"] = result["share"].cumsum()
    # Include the first item that crosses the 80% threshold.
    before = result["cumulative_share"] - result["share"]
    result["priority"] = (before < 0.8) & (result["cost"] > 0)
    return result


def export_workbook(frame: pd.DataFrame) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Summary"
    overview.append(["Portfolio demo — synthetic data"])
    overview.append(["All records and prices are fictional. Currency: illustrative USD."])
    overview.append(["Selection", "Value"])
    overview.append(["Clients", ", ".join(sorted(frame["client"].unique())) or "None"])
    overview.append(["Months", ", ".join(sorted(frame["month"].unique())) or "None"])
    for label, value in metrics(frame).items():
        overview.append([label, value])
    overview.append(["Cost per kg calculation", "Total selected cost / unique batch kilograms"])
    overview.append(["Pareto priority", "Smallest cost-ranked set reaching at least 80%"])
    for name, table in [
        ("Monthly", monthly_summary(frame)),
        ("Batches", batch_summary(frame)),
        ("Material costs", material_pareto(frame)),
        ("Detail", frame),
    ]:
        sheet = workbook.create_sheet(name)
        sheet.append(list(table.columns))
        for row in table.itertuples(index=False, name=None):
            sheet.append([value.to_pydatetime() if isinstance(value, pd.Timestamp) else value for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in workbook:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324D")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.is_date:
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
    overview.column_dimensions["A"].width = 35
    overview.column_dimensions["B"].width = 80
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
