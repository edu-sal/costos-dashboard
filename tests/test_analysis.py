import io
from pathlib import Path
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

from analysis import batch_summary, export_workbook, filter_records, material_pareto, metrics, monthly_summary, read_records
from demo_data import COLUMNS, generate_workbook


def small_frame():
    rows = [
        ["B1", "2026-05-02", "Client A", 100.0, 5, "Material A", "Materials", "kg", 5.0, 11.0, 55.0],
        ["B1", "2026-05-02", "Client A", 100.0, 5, "Material B", "Materials", "kg", 3.0, 10.0, 30.0],
        ["B2", "2026-06-02", "Client B", 300.0, 12, "Water", "Utilities", "L", 15.0, 1.0, 15.0],
    ]
    frame = pd.DataFrame(rows, columns=COLUMNS)
    frame["date"] = pd.to_datetime(frame["date"])
    frame["month"] = frame["date"].dt.strftime("%Y-%m")
    return frame


class AnalysisTests(unittest.TestCase):
    def test_weights_are_not_duplicated_and_ratio_is_weighted(self):
        frame = small_frame()
        totals = metrics(frame)
        self.assertEqual(totals["batches"], 2)
        self.assertEqual(totals["base_kg"], 400)
        self.assertEqual(totals["total_cost"], 100)
        self.assertEqual(totals["cost_per_kg"], .25)
        self.assertEqual(batch_summary(frame).set_index("batch_id").loc["B1", "cost"], 85)

    def test_filters_and_months(self):
        selected = filter_records(small_frame(), ["Client A"], ["2026-05"])
        self.assertEqual(len(selected), 2)
        monthly = monthly_summary(selected)
        self.assertEqual(monthly.iloc[0]["base_kg"], 100)
        self.assertEqual(monthly.iloc[0]["cost"], 85)
        self.assertEqual(monthly.iloc[0]["batches"], 1)
        self.assertEqual(metrics(filter_records(small_frame(), ["Unknown"]))["total_cost"], 0)

    def test_pareto_includes_the_item_crossing_eighty_percent(self):
        pareto = material_pareto(small_frame())
        self.assertEqual(pareto["priority"].tolist(), [True, True, False])
        self.assertAlmostEqual(pareto["cumulative_share"].iloc[1], .85)
        self.assertAlmostEqual(pareto["cumulative_share"].iloc[-1], 1.0)
        self.assertEqual(material_pareto(small_frame().iloc[:1])["priority"].tolist(), [True])

    def test_generated_excel_can_be_parsed_and_checked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = generate_workbook(Path(directory) / "synthetic.xlsx")
            frame = read_records(path)
            self.assertEqual(len(frame), 432)
            self.assertEqual(frame["batch_id"].nunique(), 72)
            self.assertEqual(set(frame["client"]), {"Client A", "Client B", "Client C"})
            workbook = load_workbook(path)
            workbook["Records"]["K2"] = 999999
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "Line costs"):
                read_records(path)

    def test_export_contains_only_filtered_rows_and_reconciles(self):
        frame = filter_records(small_frame(), ["Client A"])
        workbook = load_workbook(io.BytesIO(export_workbook(frame)), data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Summary", "Monthly", "Batches", "Material costs", "Detail"])
            detail = list(workbook["Detail"].values)
            headers = list(detail[0])
            self.assertEqual(len(detail) - 1, 2)
            self.assertEqual({row[headers.index("client")] for row in detail[1:]}, {"Client A"})
            self.assertEqual(sum(row[headers.index("line_cost")] for row in detail[1:]), 85)
            batches = list(workbook["Batches"].values)
            self.assertEqual(len(batches) - 1, 1)
            self.assertEqual(batches[1][list(batches[0]).index("base_kg")], 100)
            summary = {row[0]: row[1] for row in list(workbook["Summary"].values)[2:] if row[0]}
            self.assertEqual(summary["total_cost"], 85)
            self.assertEqual(summary["base_kg"], 100)
            self.assertIn("synthetic", workbook["Summary"]["A1"].value)
        finally:
            workbook.close()

    def test_empty_selection_exports_headers_without_rows(self):
        workbook = load_workbook(io.BytesIO(export_workbook(small_frame().iloc[0:0])), data_only=True)
        try:
            self.assertEqual(workbook["Detail"].max_row, 1)
            self.assertEqual(workbook["Monthly"].max_row, 1)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
